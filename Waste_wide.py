from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from config import orders_wide
from utils.Xls_row import XlsRow
from utils.core import PrettyProperties

class Waste_wide_row(XlsRow):
    def __init__(self, ws: Worksheet, row:int):
        self.ws = ws
        self.row = row

    @property
    def actual(self):
        if self.strike(2) or self.strike(3) or self.__contains__(['втратив']):
            return False

    @property
    def status_description(self):
        if not isinstance(self.cell(25).value, str):
            return ''
        return self.cell(25).value

    @property
    def is_actual(self):
        return False

    @property
    def date(self):
        return '2000-00-00'

    @property
    def ordernum(self):
        return 0000

if __name__=="__main__":
    wb = load_workbook(orders_wide, data_only=True)
    ws = wb['Sheet1']
    r = Waste_wide_row(ws, 9436)
    print(r.strike(2), r.status_description)
