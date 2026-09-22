from enum import Enum

from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from utils.core import remove_duplicate_spaces

sheets = ['РАО', 'ЗББ та Р', 'Зас ураж', 'НСО', 'БПЛА', 'ППО', 'ОВТ та МСП', 'РЕБ', 'Інж', 'РХБЗ', 'Реч', 'Звяз', 'Прод', 'ПММ', 'Мед', 'Авто', 'КЕС', 'Елек', 'Пожежна', 'Засоби розвідки', 'Гео', 'Метр']
sheets_normal = [ 'ЗББ та Р','Зас ураж','НСО','БПЛА', 'ППО', 'ОВТ та МСП', 'РЕБ', 'Інж', 'РХБЗ', 'Реч', 'Звяз', 'Прод', 'ПММ', 'Мед', 'Авто', 'КЕС', 'Елек', 'Пожежна', 'Засоби розвідки', 'Гео', ]
sheets_perversed=['РАО','Метр']

def get_status_for_row(wb: Workbook, sheetname, row):
    if sheetname not in sheets:
        raise Exception('cant find a sheet {}'.format(sheetname))
    if not wb.data_only:
        raise Exception('unmaintainable file mode')

    ws = wb[sheetname]

class Waste_status(Enum):
    BLANK = 'blank'
    CANCELED = "canceled"
    UPDATED = "updated"
    WROTEOFF = "wroteoff"
    GUILTED = "guilted"
    INQUIRY = "inquiry"
    ETC = "etc"

def waste_status(value:str):
    if value is None:
        return Waste_status.BLANK
    if not isinstance(value, str):
        return Waste_status.ETC
    value = remove_duplicate_spaces(value.strip().lower())
    if value.startswith("єас №"):
        return Waste_status.WROTEOFF
    elif value.startswith("зміни") or  value.startswith("змінено"):
        return Waste_status.UPDATED
    elif value.startswith("скасування") or "втрату" in value or "втратив" in value:
        return Waste_status.CANCELED
    elif "розслідування" in value or "рослідування" in value or "розслідуваня" in value:
        return Waste_status.INQUIRY
    elif ("винн" in value) or ("стягнути" in value):
        return Waste_status.GUILTED
    elif len(value) > 1:
        return Waste_status.ETC
    else:
        return Waste_status.BLANK

def is_guilty_person(value):
    status = waste_status(value)
    if value == Waste_status.GUILTED:
        return True
    return False

def canceled_or_updated(ws: Worksheet, row: int):
    if (ws.cell(row=row, column=2).font.strike or ws.cell(row=row, column=9).font.strike):
        return True
    if isinstance(ws.cell(row=row, column=16).value,str):
        status = ws.cell(row=row, column=16).value.lower()
        if ('скасован' in status) or ('втратив чинність' in status):
            return True

def get_wastetitle(ws: Worksheet, row: int):
    return TitleWastebook(str=ws.cell(row=row, column=4).value, total_for_items=ws.cell(row=row, column=8).value)

def get_wasteorder(ws: Worksheet, row: int):
    return OrderWasteAccountBook(date=cell_to_sqlite_date(ws.cell(row=row, column=2)))

def check_spoiled_orders(ws: Worksheet):
    all_present = True
    last_row = ws.max_row  # 8000 #9475 6078 #
    row = 12
    spoiled_rows = []
    for row in range(row, last_row):
        striked = True if (ws.cell(row=row, column=2).font.strike or ws.cell(row=row, column=9).font.strike) else False
        if canceled_or_updated(ws, row) or is_guilty_person(ws.cell(row=row, column=16).value):
            continue
        elif isinstance(ws.cell(row=row, column=9).value, float):
            year = None
            for col in range(10, 15):
                if isinstance(ws.cell(row=row, column=col).value, float):
                    year = 2012 + col
                    #print(row, col, year)
                    break
            if year is None:
                print(' no year for ', row,ws.cell(row=row, column=16).value)
                all_present = False
                spoiled_rows.append(row)
            else:

                pass #ws.cell(row=row, column=1, value=year)

    return spoiled_rows

def year_servnumb_forrao(col):
    ch,zal = divmod(col-10, 3)
    year = 2022 + zal
    return year, ch+1

def check_spoiled_orders_forrao(ws):
    all_present = True
    last_row=2983 # ws.max_row  # 8000 #9475 6078 #
    row = 12
    spoiled_rows = []
    for row in range(row, last_row):
        striked = True if (ws.cell(row=row, column=2).font.strike or ws.cell(row=row, column=9).font.strike) else False
        if canceled_or_updated(ws, row) or is_guilty_person(ws.cell(row=row, column=16).value):
            continue
        elif isinstance(ws.cell(row=row, column=9).value, float):
            year = None
            for col in range(10, 19):
                if isinstance(ws.cell(row=row, column=col).value, float):
                    year, serv = year_servnumb_forrao(col)
                    # print(row, col, year)
                    break
            if year is None:
                print(' no year for ', row, ws.cell(row=row, column=19).value)
                all_present = False
                spoiled_rows.append(row)
            else:
                pass  # ws.cell(row=row, column=1, value=year)

    return spoiled_rows

def merged(ws: Worksheet, cell = "D6"):
    for merged_range in ws.merged_cells.ranges:
        if cell in merged_range:
            coords = merged_range.start_cell.coordinate
            print("Головна комірка:", coords)
            value = ws[coords].value
            print(ws[coords].value)
            print(merged_range.min_row)
            print(merged_range.max_row)
            return

def is_actual_row(ws: Worksheet, row):
    d = cell_to_sqlite_date(ws.cell(row=row, column=2))
    n = get_order_from_comment(ws.cell(row=row, column=3).value)
    if d is None or n is None:
        #print('no d or n')
        return False
    s = ws.cell(row=row, column=2).font.strike
    if s == True:
        return False
    return True

def get_total(ws: Worksheet, row):
    if not isinstance(ws.cell(row=row, column=9).value, float):
        return None
    return Decimal(ws.cell(row=row, column=9).value)

def get_order_year(ws: Worksheet, row: int):
    if canceled_or_updated(ws, row) or is_guilty_person(ws.cell(row=row, column=16).value):
        return None
    if isinstance(ws.cell(row=row, column=9).value, float):
        year = None
        for col in range(10, 15):
            if isinstance(ws.cell(row=row, column=col).value, float):
                year = 2012 + col
                # print(row, col, year)
                break
        if year is None:
            print(' no year for ', row, ws.cell(row=row, column=16).value)
            all_present = False
        else:
            return year # ws.cell(row=row, column=1, value=year)


if __name__=='__main__':
    print(waste_status('вважати таким що втратив чинність на підставі наказу №907 від 10.02.2026').value)

#стягнути з Бухнацевич Д.В.
