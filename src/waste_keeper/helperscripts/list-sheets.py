from openpyxl import load_workbook

# Шлях до Excel-файлу
from waste_keeper.config import tmp_shortage_ledger # "data.xlsx"
wb = load_workbook(tmp_shortage_ledger)

# Вивести список усіх листів
print("Листи у файлі:")
list = []

for i, sheet_name in enumerate(wb.sheetnames, start=1):
    list.append(sheet_name)

print(list)