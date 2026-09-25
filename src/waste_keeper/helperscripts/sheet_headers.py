from openpyxl.reader.excel import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from waste_keeper.config import tmp_property_loss_orders
department_headers = ['РАО', 'РАО ЗББ та Р', 'Зас УРАЖ', 'БПЛА', 'ППО', 'НСО', 'РЕБ', 'ОВТ та МСП', 'реч', 'інж', 'зв', 'РХБЗ', 'АС', 'прод', 'мед', 'ПММ', 'гео', 'кес', 'ел-тех', 'пожежна', 'Інженерна', 'засоби розвідки', 'метрол']


if __name__ == "__main__":
    wb = load_workbook(tmp_property_loss_orders)
    ws = wb['Sheet1']
    headers = [cell.value for cell in ws[1]]
    print(headers)