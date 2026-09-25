from datetime import datetime, date
import re
from typing import Optional, Dict
from decimal import Decimal, InvalidOperation
from typing import Any
import math
import hashlib
import inspect
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils.datetime import from_excel
from enum import Enum
import shutil
from datetime import date

from waste_keeper.config import sample_xlsx
from waste_keeper.utils.row_marker import mark_row_wcolor, Warning_color
from waste_keeper.config import property_loss_orders_readonly, tmp_property_loss_orders, shortage_ledger_readonly, tmp_shortage_ledger

# sourcefile = "накази_втрати майна  А4007.xlsx"
# outputfile = "книга втрат електронний варіант.xlsx"

class PrettyProperties:
    def __repr__(self):
        props = []
        for name, _ in inspect.getmembers(type(self), lambda x: isinstance(x, property)):
            try:
                value = getattr(self, name)
            except Exception as e:
                value = f"<Error: {e}>"
            props.append(f"{name}={value!r}")
        return f"{type(self).__name__}({', '.join(props)})"

def print_properties(obj):
    props = inspect.getmembers(type(obj), lambda x: isinstance(x, property))
    if not props:
        print("No @property attributes")
        return
    width = max(len(name) for name, _ in props)
    for name, _ in props:
        try:
            value = getattr(obj, name)
        except Exception as e:
            value = f"<Error: {e}>"
        print(f"{name:<{width}} : {value}")

def remove_duplicate_spaces(text: str) -> str:
    return " ".join(text.split())

def mark_rows_from_dict(filepath, dict, color=Warning_color.GENERAL_CASE):#filepath, dict: Dict
    dict = {
        "БПЛА": [3],
    }
    wb = load_workbook(filepath, data_only=False)

    for sheetname, rows in dict.items():
        print(sheetname, rows)
        ws: Worksheet = wb[sheetname]
        for row in rows:
            print(row)
            mark_row_wcolor(ws, row, color)
    wb.save(filepath)
    wb.close()

def is_valid_date(value):
    if value is None:
        return False

    if isinstance(value, (datetime, date)):
        return True

    if not isinstance(value, str):
        return False

    value = value.strip()
    if not value:
        return False

    formats = (
        "%d.%m.%Y",
        "%d.%m.%y",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
    )

    for fmt in formats:
        try:
            datetime.strptime(value, fmt)
            return True
        except ValueError:
            pass

    return False

def cell_to_sqlite_date(value) -> str | None:
    """
    Перетворює будь-яке значення в дату формату YYYY-MM-DD.

    Повертає:
        '2024-09-02' або None.
    """

    # порожня ячейка
    if value is None:
        return None
    # datetime
    if isinstance(value, datetime):
        return value.date().isoformat()
    # date
    if isinstance(value, date):
        return value.isoformat()

    # Excel serial date (число)
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date().isoformat()
        except Exception:
            return None

    # рядок
    if isinstance(value, str):

        value = value.strip()

        if not value:
            return None

        formats = (
            "%d.%m.%Y",
            "%d/%m/%Y",
            "%Y-%m-%d",
            "%d-%m-%Y",
            "%d/%m/%y",
        )

        for fmt in formats:
            try:
                return datetime.strptime(value, fmt).date().isoformat()
            except ValueError:
                pass

        return None

    return None

def value_to_sqlite_date(value) -> str | None:
    """
    Перетворює будь-яке значення в дату формату YYYY-MM-DD.

    Повертає:
        '2024-09-02' або None.
    """
    # порожня ячейка
    if value is None:
        return None
    # datetime
    if isinstance(value, datetime):
        return value.date().isoformat()
    # date
    if isinstance(value, date):
        return value.isoformat()

    # Excel serial date (число)
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date().isoformat()
        except Exception:
            return None

    # рядок
    if isinstance(value, str):

        value = value.strip()

        if not value:
            return None

        formats = (
            "%d.%m.%Y",
            "%d/%m/%Y",
            "%Y-%m-%d",
            "%d-%m-%Y",
            '%d/%m/%y',
        )

        for fmt in formats:
            try:
                return datetime.strptime(value, fmt).date().isoformat()
            except ValueError:
                pass

        return None

    return None

def value_to_pydate(value)->date:
    return date.fromisoformat(value_to_sqlite_date(value))

def safe_decimal(value: Any) -> Decimal:
    """
    Безпечно перетворює значення на Decimal.

    Повертає:
        Decimal(...) - якщо перетворення успішне;
        Decimal(0)   - якщо перетворення неможливе.
    """
    try:
        if value is None:
            return Decimal(0)

        if isinstance(value, Decimal):
            return value

        if isinstance(value, int):
            return Decimal(value)

        if isinstance(value, float):
            # Через str(), щоб уникнути похибки двійкового представлення float
            return Decimal(str(value))

        text = str(value).strip()
        if not text:
            return Decimal(0)

        # Прибираємо пробіли між цифрами (151 150.40 -> 151150.40)
        text = text.replace(" ", "")

        return Decimal(text)

    except (InvalidOperation, ValueError, TypeError):
        return Decimal(0)

def get_order_from_comment(value) -> Optional[int]:
    value = _get_order_from_comment(value)
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        if not value or value.lower() == "none":
            return None

    try:
        return int(value)
    except (TypeError, ValueError):
        return None

_NUMBER_RE = re.compile(r"(\d+)\s*$")
def _get_order_from_comment(s) -> Optional[int]:
    if s is None:
        return None

    operation = str(s).strip()
    if not operation:
        return None

    operation = re.sub(r"\([^)]*\)", "", operation).strip()

    match = _NUMBER_RE.search(operation)
    if match:
        return int(match.group(1))

    return None

def cell_to_decimal(cell) -> Decimal | None:
    """
    Безпечно перетворює значення openpyxl.cell.Cell у Decimal.

    Повертає:
        Decimal - якщо значення коректне.
        None    - якщо значення відсутнє або некоректне.
    """
    value = cell.value

    if value is None:
        return None

    # Уже Decimal
    if isinstance(value, Decimal):
        return value

    # Цілі числа
    if isinstance(value, int):
        return Decimal(value)

    # float з Excel
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        # Через str(), щоб не тягнути двійкові похибки float
        return Decimal(str(value))

    # Рядок
    if isinstance(value, str):
        s = (
            value.strip()
            .replace("\xa0", "")   # нерозривний пробіл
            .replace(" ", "")      # звичайний пробіл
            .replace(",", ".")     # десяткова кома -> крапка
        )

        if not s:
            return None

        try:
            return Decimal(s)
        except InvalidOperation:
            return None

    # Інші типи
    return None

def calculate_md5(file_path) -> str:
    md5 = hashlib.md5()
    with open(file_path, "rb") as f:
        while chunk := f.read(1024 * 1024):
            md5.update(chunk)
    return md5.hexdigest()

def is_valid_order_row(ws, row):
    if ws.cell(row=row, column=2).value is None or ws.cell(row=row, column=3).value is None:
        return False
    textdate = ws.cell(row=row, column=2).value
    if not is_valid_date(textdate):
        print('cant get order or data', row, textdate)
        return False
    return True

def get_order_from_comment_old(s="(зміни в 2431)    3719"):
    operation = str(s).strip()
    if operation.isdigit():
        order_id = int(operation)
    else:
        order_id = re.sub(r'\([^)]*\)', '', operation).strip()
        order_id = order_id.split()[-1]
    return order_id

def remove_last_number(s="(зміни в 2431)                               3719"):
    if s is None:
        return None
    s=str(s)
    m = re.search(r'^(.*?)([+-]?\d+)\s*$', s)
    return m.group(1).rstrip() if m else s

def format_date_for_output(date_str):
    return datetime.strptime(date_str, "%Y-%m-%d").strftime("%d.%m.%Y")

def delete_empty_bottom(ws:Worksheet):
    start_row = None
    empty_count = 0

    for row in range(ws.max_row, 0, -1):
        is_empty = all(
            cell.value in (None, '')
            for cell in ws[row]
        )

        if is_empty:
            start_row = row
            empty_count += 1
        else:
            break

    if empty_count > 0:
        print('deleting ', empty_count)
        ws.delete_rows(start_row, empty_count)
    else:
        print('no empty rows')

def safe_int(value: str) -> int | None:
    if not isinstance(value, str):
        return None

    value = value.lstrip()

    digits = ""
    for char in value:
        if not char.isdigit():
            break
        digits += char

    return int(digits) if digits else None

def grap_operable_property_loss_orders():
    shutil.copy2(property_loss_orders_readonly, tmp_property_loss_orders)

def get_last_excel_files():
    shutil.copy2(property_loss_orders_readonly, tmp_property_loss_orders)
    shutil.copy2(shortage_ledger_readonly, tmp_shortage_ledger)

#date oh headers 23/09/2026
headers = ['РАО', 'РАО ЗББ та Р', 'Зас УРАЖ', 'БПЛА', 'ППО', 'НСО', 'РЕБ', 'ОВТ та МСП', 'реч', 'інж', 'зв', 'РХБЗ', 'АС', 'прод', 'мед', 'ПММ', 'гео', 'кес', 'ел-тех', 'пожежна', 'Інженерна', 'засоби розвідки', 'метрол']

#date oh sheet names 24/09/2026
sheets = ['ЗББ та Р', 'Зас ураж', 'НСО', 'БПЛА', 'ППО', 'ОВТ та МСП', 'РЕБ', 'Реч', 'Звяз', 'Прод', 'ПММ', 'Мед', 'Авто', 'КЕС', 'Елек', 'Інженерна', 'Пожежна', 'Засоби розвідки', 'Гео', 'Метр', 'РАО', 'РХБЗ', 'Інж']


department_to_sheet = {
    None: None,
    "рао": None,
    "рао збб та р": "ЗББ",
    "зас ураж": "ЗУ",
    "бпла": "БпЛА",
    "ппо": None,
    "нсо": "НСО (БТ)",
    "реб": None,
    "овт та мсп": "ОВТ",
    "реч": "РЕЧ",
    "інж": "ІС",
    "зв": "ЗВ",
    "рхбз": None,
    "ас": "СВТ (АС)",
    "прод": "прод",
    "мед": "мед",
    "пмм": "пмм",
    "гео": "ГЕО",
    "кес": "КЕС( СІ-ІЗ)",
    "елтех": "Ел-тех",
    "пожежна": None,
    "метрол": "Метрологія",
}

if __name__=='__main__':
    get_last_excel_files()
    # dict = {
    #     "БПЛА": [3],
    # }
    # mark_rows_from_dict(sample_xlsx, dict)
