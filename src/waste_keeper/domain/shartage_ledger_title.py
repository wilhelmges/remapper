from openpyxl import load_workbook

from waste_keeper.domain.order_meta_extractor import OrderMetaExtractor


class ShortageLedgerTitle(OrderMetaExtractor):
    def __init__(self,ws,row):
        super().__init__(ws,row)

    @property
    def total(self)->float|None:
        if not self.isorder:
            return None
        descr = self.cell(8).value
        if isinstance(descr, float):
            return descr
        elif isinstance(descr, int):
            return float(descr)
        else:
            print(type(descr), descr)
            #raise Exception('no ideas about total')


if __name__ == '__main__':
    from waste_keeper.config import tmp_property_loss_orders, tmp_shortage_ledger
    wb = load_workbook(tmp_shortage_ledger, data_only=True)
    ws = wb["БПЛА"]

    slt = ShortageLedgerTitle(ws,16886)
    print(slt.total)
