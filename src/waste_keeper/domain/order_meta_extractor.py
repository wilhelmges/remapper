from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from waste_keeper.domain.core import (is_valid_date, value_to_sqlite_date, value_to_pydate,
                                      remove_parentheses, get_parentheses_text, safe_int)
from waste_keeper.domain.xls_row import XlsRow

class OrderMetaExtractor(XlsRow):
    def __init__(self, ws: Worksheet, row: int):
        self.ws = ws
        self.row = row


    @property
    def isorder(self):
        print(self.row)
        descr = self.ws.cell(row=self.row, column=2).value
        if descr is None:
            return False
        if is_valid_date(descr):
            return True
        return False

    @property
    def isactive(self):
        if self.ws.cell(row=self.row, column=2).font.strike or self.ws.cell(row=self.row, column=3).font.strike:
            return False
        return True

    @property
    def orderdate(self):
        if not self.isorder:
            return None
        return value_to_pydate(self.ws.cell(row=self.row,column=2).value)

    @property
    def ordernum(self)->int|None:
        if not self.isorder:
            return None
        descr = self.ws.cell(row=self.row, column=3).value

        if descr is None:
            return None
        elif isinstance(descr, int):
            return descr
        elif isinstance(descr, str):
            descr = descr.strip().lower()
            descr = remove_parentheses(descr)
            descr = ''.join(char for char in descr if not char.isalpha())

            #els = (descr.split(' '))
            return safe_int(descr)
        else:
            print(type(descr))
            raise Exception(f'what to do with {descr}')


if __name__=="__main__":
    from waste_keeper.config import tmp_property_loss_orders, tmp_shortage_ledger
    wb: Workbook = load_workbook(tmp_shortage_ledger)
    ws: Worksheet = wb["БПЛА"]

    ome: OrderMetaExtractor = OrderMetaExtractor(ws, 4056)
    print(ome.isorder,ome.isactive, ome.orderdate, ome.ordernum)
    print(ome)
