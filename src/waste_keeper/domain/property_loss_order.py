from openpyxl.reader.excel import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from functools import wraps
import re

from waste_keeper.utils.core import is_valid_date, value_to_pydate
from waste_keeper.utils.core import safe_int

def is_special(func):
    """check special case for order status"""
    @wraps(func)
    def wrapper(self):
        if self.canceled and self.description is None:
            return None
        if isinstance(self.description, str):
            if not self.description.startswith("втратив") and not self.description.startswith("зміни"):
                return None

        return func(self)
    return wrapper

def remove_parentheses(text):
    result = []
    depth = 0

    for char in text:
        if char == '(':
            depth += 1
        elif char == ')':
            if depth > 0:
                depth -= 1
        elif depth == 0:
            result.append(char)

    return ' '.join(''.join(result).split())

def get_parentheses_text(s: str) -> str | None:
    start = s.find("(")
    if start == -1:
        return None

    level = 0

    for i in range(start, len(s)):
        if s[i] == "(":
            level += 1
        elif s[i] == ")":
            level -= 1

            if level == 0:
                return s[start:i + 1]

    return None

class Property_loss_order:
    def __init__(self, ws: Worksheet, row:int):
        self.ws: Worksheet = ws
        self.row: int = row

    @property
    def isorder(self):
        descr = self.ws.cell(row=self.row, column=2).value
        if descr is None:
            return False
        if is_valid_date(descr):
            return True
        return False

    @property
    def orderdate(self):
        if not self.isorder:
            return None
        return value_to_pydate(self.ws.cell(row=self.row,column=2).value)

    @property
    def ordernum(self)->int|None:
        if not self.isorder:
            return None
        descr = self.ws.cell(row=self.row, column=3).value

        if descr is None:
            return None
        elif isinstance(descr, int):
            return descr
        elif isinstance(descr, str):
            descr = descr.strip().lower()
            descr = remove_parentheses(descr)
            descr = ''.join(char for char in descr if not char.isalpha())

            #els = (descr.split(' '))
            return safe_int(descr)
        else:
            print(type(descr))
            raise Exception(f'what to do with {descr}')

    @property
    def description(self):
        value = self.ws.cell(row=self.row,column=27).value
        if value is None:
            value = self.ws.cell(row=self.row, column=28).value
        if isinstance(value, str):
            return value.strip().lower()
        return value

    @property
    def canceled(self):
        if self.canceled_by_strike:
            return True
        descr = self.description
        if descr is None:
            return False
        if isinstance(descr, str):
            descr = descr.strip().lower()
            return descr.startswith("втратив")
        return False

    @property
    def canceled_by_strike(self):
        if self.ws.cell(row=self.row,column=2).font.strike or self.ws.cell(row=self.row,column=3).font.strike:
            return True

    @property
    @is_special
    def canceled_by_order(self)->str|None:
        if not self.canceled:
            return None
        descr = self.description
        descr = remove_parentheses(descr)
        elements = descr.split("№")
        return elements[-1]

    @property
    @is_special
    def canceled_by_ordernum(self):
        if not self.canceled:
            return None
        descr = self.canceled_by_order
        els = descr.split(" ")
        return int(els[0])

    @property
    @is_special
    def canceled_by_orderdate(self):
        if not self.canceled:
            return None
        descr = self.canceled_by_order[-10:]
        descr = ''.join(char for char in descr if not char.isalpha())
        return value_to_pydate(descr)

    @property
    def effects_to_order_num(self):
        if not self.isorder or self.canceled:
            return None
        descr = self.ws.cell(row=self.row,column=3).value
        if isinstance(descr, int):
            return None
        return get_parentheses_text(descr)


if __name__=='__main__':
    from waste_keeper.config import tmp_property_loss_orders
    wb: Workbook = load_workbook(tmp_property_loss_orders, data_only=True)
    ws: Worksheet = wb["Sheet1"]
    plo: Property_loss_order = Property_loss_order(ws, 7119)
    print(plo.ordernum, plo.effects_to_order_num)
