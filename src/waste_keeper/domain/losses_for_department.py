from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


def get_list(ws, num):
    list = []
    for column in range(4, 27):
        descr = ws.cell(row=num, column=column).value
        if descr is None:
            list.append(None)
        elif isinstance(descr, float):
            list.append((descr))
        elif isinstance(descr, int):
            list.append(float(descr))
        elif isinstance(descr, str):
            descr = descr.strip().replace(',', '.')
            try:
                value = float(descr)
                list.append(value)
            except Exception:
                list.append(None)
        else:
            raise Exception(f"type descr: {type(descr)}")
    return list


if __name__ == '__main__':
    from waste_keeper.config import tmp_property_loss_orders
    wb: Workbook = load_workbook(tmp_property_loss_orders, data_only=True)
    ws: Worksheet = wb["Sheet1"]
    print(get_list(ws, 14394))