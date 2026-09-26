from openpyxl import load_workbook

from waste_keeper.domain.shartage_ledger_title import ShortageLedgerTitle


class Shortage_Ledger_group(ShortageLedgerTitle):
    def __init__(self, ws, row):
        super().__init__(ws, row)

    @property
    def isgroup(self):
        if not self.isorder:
            return False
        return True

    @property
    def startnum(self):
        if not self.isgroup:
            return None
        num = self.row
        ordernum = self.ordernum
        while True:
            slt = ShortageLedgerTitle(self.ws, num-1)
            if slt.ordernum != ordernum:
                break
            num -= 1
        return num

    @property
    def endnum(self):
        if not self.isgroup:
            return None
        num = self.row
        ordernum = self.ordernum
        while True:
            slt = ShortageLedgerTitle(self.ws, num+1)
            if slt.ordernum != ordernum:
                break
            num += 1
        return num

    @property
    def description(self):
        slt: ShortageLedgerTitle = ShortageLedgerTitle(self.ws, self.startnum)
        descr = slt.cell(16).value
        if descr is not None:
            return descr
        slt: ShortageLedgerTitle = ShortageLedgerTitle(self.ws, self.endnum)
        return slt.cell(16).value


    @property
    def total_amount(self):
        raise Exception('not implemented')


if __name__ == '__main__':
    from waste_keeper.config import tmp_property_loss_orders, tmp_shortage_ledger
    wb = load_workbook(tmp_shortage_ledger, data_only=True)
    ws = wb["БПЛА"]

    slg = Shortage_Ledger_group(ws, 4315)
    print(slg.startnum, slg.ordernum, slg.description)


