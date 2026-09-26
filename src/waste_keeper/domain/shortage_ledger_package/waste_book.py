from pathlib import Path
from dataclasses import dataclass
from decimal import Decimal
import datetime
from types import SimpleNamespace

from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from utils.core import cell_to_sqlite_date, get_order_from_comment, cell_to_decimal, safe_decimal
from utils.row_marker import Warning_color, mark_row_wcolor
from config import wasted_backup as file_path, sample_xlsx, wasted_backup  # "data.xlsx"
from row_factory import row_factory
from waste_core import sheets_normal

from sqlmodel import Session
from database.mysqlmodel import engine
from database.models import Test2, OrderWasteAccountBook, TitleWastebook
from database.Repository import Repository
from database.utils import native_clear_table_rows

sheets = ['РАО', 'ЗББ та Р', 'Зас ураж', 'НСО', 'БПЛА', 'ППО', 'ОВТ та МСП', 'РЕБ', 'Інж', 'РХБЗ', 'Реч', 'Звяз', 'Прод', 'ПММ', 'Мед', 'Авто', 'КЕС', 'Елек', 'Пожежна', 'Засоби розвідки', 'Гео', 'Метр']

def parse_by_sheetname(wb: Workbook, sheetname: str):
    ws = wb[sheetname]
    uni_waste_row = row_factory(sheetname)

    with Session(engine) as session:
        orepo = Repository(OrderWasteAccountBook, session)
        trepo = Repository(TitleWastebook, session)

        last_row = ws.max_row  # 8000 #9475 6078 #
        first_row =  13170 #7
        previous_order_num = None; set_total_sum = 0
        order_num = None
        for row in range(last_row,first_row,-1):
            wr = uni_waste_row(wb, sheet_name, row)
            if wr.is_datarow:
                print(order_num, str(wr.status.value))

                if wr.ordernum is not None and wr.ordernum != order_num:
                    order_num = wr.ordernum
                    order = orepo.create(OrderWasteAccountBook(
                        ordernum=wr.ordernum,date = wr.sqldate, totalsum=wr.total_sum,
                        status=wr.status.value, status_description=wr.status_description, year=wr.year,
                        eventyear = wr.year, strike=wr.is_striken, unit = sheetname,
                        sheetrow = row
                    ))
                else:
                    pass
                item = trepo.create(TitleWastebook(title=wr.title,amount=wr.amount, sheetrow=row, order_id=order.id))

if __name__ == '__main__':
    wb = load_workbook(wasted_backup, data_only=True)
    sheet_name = 'БПЛА'
    ws = wb[sheet_name]
    native_clear_table_rows(['orders_wasteaccountbook', 'titles_wastebook'])

    parse_by_sheetname(wb, sheet_name); exit(0)

    for sheet_name in sheets_normal:
        ws = wb[sheet_name]
        print(sheet_name)
        parse_by_sheetname(wb, sheet_name);
    exit(0)


    exit(0)


def parse_sheet_legacy(ws: Worksheet):
    last_row =  ws.max_row # 8000 #9475 6078 #
    row = 7
    previous_order_num = None; set_total_sum = 0
    for row in range(row, last_row):
        date = cell_to_sqlite_date(ws.cell(row=row, column=2))
        ordernum = (get_order_from_comment(ws.cell(row=row, column=3).value))

        if (date is None or ordernum is None) and previous_order_num is None:
            print(row, 'empty row')
            continue

        # new set started
        elif ((date and ordernum) and (previous_order_num != ordernum)) or ((date is None or ordernum is None) and previous_order_num is not None):
            if previous_order_num is not None:
                set_total_sum2 = ( cell_to_decimal(ws.cell(row=row-1, column=9)))
                if set_total_sum2 is not None and abs(abs(set_total_sum) - abs(set_total_sum2))>0.001:
                    mark_row_wcolor(ws, row-1, Warning_color.DIFF_SUMS)
                    print('does not equal for ', row-1, set_total_sum, set_total_sum2)
                else:
                    pass #print(set_total_sum, set_total_sum2)
            #print('------------------------------')
            #print('new set started with ordernum ', ordernum, type(ordernum))
            previous_order_num = ordernum
            set_total_sum = Decimal(0)
            changed = True if ws.cell(row=row, column=2).font.strike else False
            title = ws.cell(row=row, column=4).value
            title_total_sum = safe_decimal(ws.cell(row=row, column=8).value)
            set_total_sum += title_total_sum

            # print(row, date, ordernum, changed, title, title_total_sum)

        elif (date and ordernum) and (previous_order_num == ordernum and previous_order_num):
            # print('set continued', row, ordernum, type(ordernum))
            changed = True if ws.cell(row=row, column=2).font.strike else False
            title = ws.cell(row=row, column=4).value
            title_total_sum = safe_decimal(ws.cell(row=row, column=8).value)
            set_total_sum+=title_total_sum

            #print(row, date, ordernum, changed, title, title_total_sum)

            if title_total_sum is None:
                pass #print('file may be need to recalculate formulas'); exit(1)
                continue
        else:
            print(row, 'else')



