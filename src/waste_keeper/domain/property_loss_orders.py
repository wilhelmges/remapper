from openpyxl.reader.excel import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from waste_keeper.utils.core import is_valid_date, value_to_sqlite_date

class Property_loss_orders:
    def __init__(self, ws: Worksheet, row:int):
        self.ws: Worksheet = ws
        self.row: int = row

    @property
    def isorder(self):
        descr = self.ws.cell(row=self.row, column=2).value
        if descr is None:
            return False
        if is_valid_date(descr):
            return True
        return False

    @property
    def orderdate(self):
        if not self.isorder:
            return None
        return value_to_sqlite_date(self.ws.cell(row=self.row,column=2).value)

    @property
    def ordernum(self)->int|None:
        descr = self.ws.cell(row=self.row, column=3).value
        if not self.isorder or descr is None:
            return None

        if isinstance(descr, str):
            els = (descr.split(' '))
            return int(els[-1])

        if isinstance(descr, int):
            return descr

        print(type(descr))
        raise Exception(f'what to do with {descr}')

    @property
    def canceled(self):
        descr = self.ws.cell(row=self.row,column=27).value
        if descr is None:
            return False
        if isinstance(descr, str):
            descr = descr.strip().lower()
            return descr.startswith("втратив")
        return False

    @property
    def canceled_by_strike(self):
        if ws.cell(row=self.row,column=2).font.strike or ws.cell(row=self.row,column=3).font.strike:
            return True

    @property
    def canceled_by_order(self):
        if not self.canceled:
            return None
        descr = self.ws.cell(row=self.row,column=27).value
        elements = descr.split("№")
        return elements[1]


if __name__=='__main__':
    from waste_keeper.config import tmp_property_loss_orders
    wb = load_workbook(tmp_property_loss_orders, data_only=True)
    ws: Worksheet = wb["Sheet1"]
    plo: Property_loss_orders = Property_loss_orders(ws, 14402)
    print(plo.isorder, plo.ordernum)