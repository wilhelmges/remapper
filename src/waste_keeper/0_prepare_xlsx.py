import win32com.client
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from waste_keeper.utils.core import get_last_excel_files
from waste_keeper.config import tmp_property_loss_orders

if __name__ == "__main__":
    get_last_excel_files()

    #recalc all formulas
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    wb = excel.Workbooks.Open(tmp_property_loss_orders)
    excel.CalculateFullRebuild()
    wb.Save()
    wb.Close()
    excel.Quit()

    #patch cells
    from waste_keeper.config import tmp_property_loss_orders
    wb: Workbook = load_workbook(tmp_property_loss_orders)
    ws: Worksheet = wb["Sheet1"]

    if ws.cell(row=7095, column=27).value == 'втратив чинність наказом №1630 від 12.03.206':
        ws.cell(row=7095, column=27).value='втратив чинність наказом №1630 від 12.03.2026'

    if ws.cell(row=7008, column=27).value == 'зміни наказом №1545 від 09.03.20226':
        ws.cell(row=7008, column=27).value='зміни наказом №1545 від 09.03.2026'

    if ws.cell(row=7006, column=27).value == 'зміни наказом №1544 від 09.03.20226':
        ws.cell(row=7006, column=27).value='зміни наказом №1544 від 09.03.2026'

    if ws.cell(row=7004, column=27).value == 'зміни наказом №1543 від 09.03.20226':
        ws.cell(row=7004, column=27).value='зміни наказом №1543 від 09.03.2026'

    if ws.cell(row=7004, column=27).value == 'зміни наказом №1543 від 09.03.20226':
        ws.cell(row=7004, column=27).value='зміни наказом №1543 від 09.03.2026'

    if ws.cell(row=6979, column=27).value == 'зміни наказом 6723 від 05.11.2024':
        ws.cell(row=6979, column=27).value='зміни наказом №6723 від 05.11.2024'

    if ws.cell(row=6336, column=27).value == 'втратив чинність наказом №1630 від 12.03.206':
        ws.cell(row=6336, column=27).value='втратив чинність наказом №1630 від 12.03.2026'


    if ws.cell(row=11589, column=3).value == '(зміни в 1955                                 5717':
        ws.cell(row=11589, column=3).value='(зміни в 1955)                                5717'

    if ws.cell(row=7698, column=3).value == 'зміни в наказ (1397; 6537)          2184':
        ws.cell(row=7698, column=3).value='(зміни в наказ 1397; 6537)          2184'

    if ws.cell(row=7672, column=3).value == '(зміни в 5436(5649)                       2307':
        ws.cell(row=7672, column=3).value='(зміни в 5436(5649))                       2307'

    wb.save(tmp_property_loss_orders)
