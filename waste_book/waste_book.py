from pathlib import Path
from dataclasses import dataclass
from decimal import Decimal
import datetime
from types import SimpleNamespace

from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from utils.core import cell_to_sqlite_date, get_order_from_comment, cell_to_decimal, safe_decimal
from utils.row_marker import Warning_color, mark_row_wcolor

from database.models import TitleWastebook, OrderWasteAccountBook

from config import wasted_backup as file_path, sample_xlsx, wasted_backup  # "data.xlsx"
from Waste_row import Waste_row
from waste_core import Waste_status, waste_status

sheets = ['РАО', 'ЗББ та Р', 'Зас ураж', 'НСО', 'БПЛА', 'ППО', 'ОВТ та МСП', 'РЕБ', 'Інж', 'РХБЗ', 'Реч', 'Звяз', 'Прод', 'ПММ', 'Мед', 'Авто', 'КЕС', 'Елек', 'Пожежна', 'Засоби розвідки', 'Гео', 'Метр']
from waste_core import sheets_normal

def parse_sheet_legacy(ws: Worksheet):
    last_row =  ws.max_row # 8000 #9475 6078 #
    row = 7
    previous_order_num = None; set_total_sum = 0
    for row in range(row, last_row):
        date = cell_to_sqlite_date(ws.cell(row=row, column=2))
        ordernum = (get_order_from_comment(ws.cell(row=row, column=3).value))

        if (date is None or ordernum is None) and previous_order_num is None:
            print(row, 'empty row')
            continue

        # new set started
        elif ((date and ordernum) and (previous_order_num != ordernum)) or ((date is None or ordernum is None) and previous_order_num is not None):
            if previous_order_num is not None:
                set_total_sum2 = ( cell_to_decimal(ws.cell(row=row-1, column=9)))
                if set_total_sum2 is not None and abs(abs(set_total_sum) - abs(set_total_sum2))>0.001:
                    mark_row_wcolor(ws, row-1, Warning_color.DIFF_SUMS)
                    print('does not equal for ', row-1, set_total_sum, set_total_sum2)
                else:
                    pass #print(set_total_sum, set_total_sum2)
            #print('------------------------------')
            #print('new set started with ordernum ', ordernum, type(ordernum))
            previous_order_num = ordernum
            set_total_sum = Decimal(0)
            changed = True if ws.cell(row=row, column=2).font.strike else False
            title = ws.cell(row=row, column=4).value
            title_total_sum = safe_decimal(ws.cell(row=row, column=8).value)
            set_total_sum += title_total_sum

            # print(row, date, ordernum, changed, title, title_total_sum)

        elif (date and ordernum) and (previous_order_num == ordernum and previous_order_num):
            # print('set continued', row, ordernum, type(ordernum))
            changed = True if ws.cell(row=row, column=2).font.strike else False
            title = ws.cell(row=row, column=4).value
            title_total_sum = safe_decimal(ws.cell(row=row, column=8).value)
            set_total_sum+=title_total_sum

            #print(row, date, ordernum, changed, title, title_total_sum)

            if title_total_sum is None:
                pass #print('file may be need to recalculate formulas'); exit(1)
                continue
        else:
            print(row, 'else')

def is_guilty_person(value):
    status = waste_status(value)
    if value == Waste_status.GUILTED:
        return True
    return False

def canceled_or_updated(ws: Worksheet, row: int):
    if (ws.cell(row=row, column=2).font.strike or ws.cell(row=row, column=9).font.strike):
        return True
    if isinstance(ws.cell(row=row, column=16).value,str):
        status = ws.cell(row=row, column=16).value.lower()
        if ('скасован' in status) or ('втратив чинність' in status):
            return True

def get_wastetitle(ws: Worksheet, row: int):
    return TitleWastebook(str=ws.cell(row=row, column=4).value, total_for_items=ws.cell(row=row, column=8).value)

def get_wasteorder(ws: Worksheet, row: int):
    return OrderWasteAccountBook(date=cell_to_sqlite_date(ws.cell(row=row, column=2)))

def check_spoiled_orders(ws: Worksheet):
    all_present = True
    last_row = ws.max_row  # 8000 #9475 6078 #
    row = 12
    spoiled_rows = []
    for row in range(row, last_row):
        striked = True if (ws.cell(row=row, column=2).font.strike or ws.cell(row=row, column=9).font.strike) else False
        if canceled_or_updated(ws, row) or is_guilty_person(ws.cell(row=row, column=16).value):
            continue
        elif isinstance(ws.cell(row=row, column=9).value, float):
            year = None
            for col in range(10, 15):
                if isinstance(ws.cell(row=row, column=col).value, float):
                    year = 2012 + col
                    #print(row, col, year)
                    break
            if year is None:
                print(' no year for ', row,ws.cell(row=row, column=16).value)
                all_present = False
                spoiled_rows.append(row)
            else:

                pass #ws.cell(row=row, column=1, value=year)

    return spoiled_rows

def year_servnumb_forrao(col):
    ch,zal = divmod(col-10, 3)
    year = 2022 + zal
    return year, ch+1

def check_spoiled_orders_forrao(ws):
    all_present = True
    last_row=2983 # ws.max_row  # 8000 #9475 6078 #
    row = 12
    spoiled_rows = []
    for row in range(row, last_row):
        striked = True if (ws.cell(row=row, column=2).font.strike or ws.cell(row=row, column=9).font.strike) else False
        if canceled_or_updated(ws, row) or is_guilty_person(ws.cell(row=row, column=16).value):
            continue
        elif isinstance(ws.cell(row=row, column=9).value, float):
            year = None
            for col in range(10, 19):
                if isinstance(ws.cell(row=row, column=col).value, float):
                    year, serv = year_servnumb_forrao(col)
                    # print(row, col, year)
                    break
            if year is None:
                print(' no year for ', row, ws.cell(row=row, column=19).value)
                all_present = False
                spoiled_rows.append(row)
            else:
                pass  # ws.cell(row=row, column=1, value=year)

    return spoiled_rows

def merged(ws: Worksheet, cell = "D6"):
    for merged_range in ws.merged_cells.ranges:
        if cell in merged_range:
            coords = merged_range.start_cell.coordinate
            print("Головна комірка:", coords)
            value = ws[coords].value
            print(ws[coords].value)
            print(merged_range.min_row)
            print(merged_range.max_row)
            return

def is_actual_row(ws: Worksheet, row):
    d = cell_to_sqlite_date(ws.cell(row=row, column=2))
    n = get_order_from_comment(ws.cell(row=row, column=3).value)
    if d is None or n is None:
        #print('no d or n')
        return False
    s = ws.cell(row=row, column=2).font.strike
    if s == True:
        return False
    return True

def get_total(ws: Worksheet, row):
    if not isinstance(ws.cell(row=row, column=9).value, float):
        return None
    return Decimal(ws.cell(row=row, column=9).value)

def get_order_year(ws: Worksheet, row: int):
    if canceled_or_updated(ws, row) or is_guilty_person(ws.cell(row=row, column=16).value):
        return None
    if isinstance(ws.cell(row=row, column=9).value, float):
        year = None
        for col in range(10, 15):
            if isinstance(ws.cell(row=row, column=col).value, float):
                year = 2012 + col
                # print(row, col, year)
                break
        if year is None:
            print(' no year for ', row, ws.cell(row=row, column=16).value)
            all_present = False
        else:
            return year # ws.cell(row=row, column=1, value=year)

def parse_by_sheetname(wb: Workbook, sheetname: str):
    ws = wb[sheetname]
    last_row = ws.max_row  # 8000 #9475 6078 #
    previous_order_num = None; set_total_sum = 0
    for row in range(last_row,7,-1):
        wr = Waste_row(wb, sheet_name, row)

        order_num = None
        if wr.is_datarow:
            if wr.ordernum is not None and wr.ordernum != order_num:
                order_num =  wr.ordernum
            else:
                pass
            print(wr.cell(16).value)
        # if is_actual_row(ws, row) and not waste_row.is_guilty_person():
        #     year = get_order_year(ws, row)
        #     total_sum = get_total(ws, row)
        #     if total_sum and year is None:
        #         print('cant define year for row ', row)
        # else:
        #     pass

if __name__ == '__main__':
    wb = load_workbook(wasted_backup, data_only=True)
    sheet_name = 'Реч'
    ws = wb[sheet_name]
    parse_by_sheetname(wb, sheet_name)
    exit(0)
    wr: Waste_row = Waste_row(wb, sheet_name, 10454)
    print(wr.is_year_needed());

    for sheet_name in sheets_normal:
        ws = wb[sheet_name]
        print(sheet_name)
        parse_by_sheetname(wb, sheet_name);

    exit(0)

    last_row = ws.max_row  # 8000 #9475 6078 #
    previous_order_num = None;
    set_total_sum = 0
    for row in range(7, last_row):
        status = waste_status(ws.cell(row=row, column=16).value)
        if status == Waste_status.ETC:
            print(row, ws.cell(row=row, column=16).value)
    exit(0)

    print(is_actual_row(ws, 9936))
    print(get_order_year(ws, 4584)); exit(0)
    #print(canceled_or_updated(ws, 10454)); exit()

    for ws in wb.worksheets:
        print(ws.title)
        if ws.title in ['Count', 'Метр']:
            continue
        elif ws.title == "РАО":
            rows = check_spoiled_orders_forrao(ws)
            print(len(rows))
        else:
            rows = check_spoiled_orders(ws)
            if len(rows)>0:
                print(ws.title, rows)

    #wb.save(wasted_backup); wb.close()



