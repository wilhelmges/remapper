from openpyxl import load_workbook
from Waste_row import Waste_row
from config import wasted_backup

class Raowaste_row(Waste_row):
    @property
    def year(self):
        year = None
        for col in range(10, 15):
            if isinstance(self.cell(col).value, float):
                ch, zal = divmod(col - 10, 3)
                year = 2022 + zal
                return year #, ch + 1
        return year

if __name__ == '__main__':
    wb = load_workbook(wasted_backup, data_only=True)
    sheet_name = 'РАО'
    ws = wb[sheet_name]
    wr: Raowaste_row = Raowaste_row(wb, sheet_name, 3497)
    print(wr.year, wr.amount)
    print(wr)
