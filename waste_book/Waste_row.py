import datetime
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from config import wasted_backup
from utils.core import remove_duplicate_spaces, cell_to_sqlite_date, get_order_from_comment, PrettyProperties
from waste_core import Waste_status

class Waste_row(PrettyProperties):
    def __init__(self, wb: Workbook, sheet_name: str, row:int):
        self.wb = wb
        self.ws: Worksheet = wb[sheet_name]
        self.row = row

    def cell(self, col: int):
        return self.ws.cell(row=self.row, column=col)

    @property
    def sqldate(self):
        value = self.cell(2).value
        if value is None:
            return None
        # datetime
        if isinstance(value, datetime.datetime):
            return value.date().isoformat()
        # date
        if isinstance(value, datetime.date):
            return value.isoformat()

    @property
    def ordernum(self):
        return get_order_from_comment(self.cell(3).value)

    @property
    def status(self):
        if self.cell(2).font.strike or self.cell(9).font.strike or self.cell(16).font.strike:
            return Waste_status.CANCELED
        value = self.ws.cell(row=self.row, column=16).value
        if value is None:
            return Waste_status.BLANK
        if not isinstance(value, str):
            return Waste_status.ETC
        value = remove_duplicate_spaces(value.strip().lower())
        if value.startswith("єас №") or "інспекторське" in value :
            return Waste_status.WROTEOFF
        elif value.startswith("змінено") or "змін" in value :
            return Waste_status.UPDATED
        elif value.startswith("скасування") or "втрату" in value or "втратив" in value or "втрата" in value or "скасовано" in value:
            return Waste_status.CANCELED
        elif "розслідуван" in value or "рослідув" in value:
            return Waste_status.INQUIRY
        elif ("винн" in value) or ("стягнути" in value):
            return Waste_status.GUILTED
        elif len(value) > 1:
            return Waste_status.ETC
        else:
            return Waste_status.BLANK

    @property
    def status_description(self):
        if not isinstance(self.cell(16).value, str):
            return ''
        return self.cell(16).value

    @property
    def total_sum(self):
        if not isinstance(self.cell(9).value, float):
            return None
        return float(self.cell(9).value)

    @property
    def year(self):
            year = None
            for col in range(10, 15):
                if isinstance(self.cell(col).value, float):
                    year = 2012 + col
                    break
            return year

    #item fields
    @property
    def title(self):
        return '' + str(self.cell(4).value)

    @property
    def amount(self):
        if self.cell(8).value is None:
            return None
        return float(self.cell(8).value)

    @property
    def is_datarow(self):
        if self.sqldate and self.ordernum:
            return True

    @property
    def is_striken(self):
        status = self.status
        return (self.cell(2).font.strike or self.cell(3).font.strike
                or self.cell(4).font.strike  or self.cell(9).font.strike
                or status == Waste_status.CANCELED
                or status == Waste_status.UPDATED)

    def is_year_needed(self):
        status = self.waste_status()
        match status:
            case Waste_status.WROTEOFF | Waste_status.INQUIRY:
                return True
            case _:
                return False

    def is_actual(self):
        d = cell_to_sqlite_date(self.ws.cell(row=self.row, column=2))
        n = get_order_from_comment(self.ws.cell(row=self.row, column=3).value)
        if d is None or n is None:
            # print('no d or n')
            return False
        if self.cell(2).font.strike or self.cell(9).font.strike:
            return False
        status = self.waste_status()

        return False

if __name__=='__main__':
    wb = load_workbook(wasted_backup, data_only=True)
    sheet_name = 'БПЛА'
    ws = wb[sheet_name]
    wr: Waste_row = Waste_row(wb, sheet_name, 13182)
    print(wr.title, wr.amount)
    print(wr.is_striken)
    exit(0)